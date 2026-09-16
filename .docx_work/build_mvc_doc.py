from __future__ import annotations

import os
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\workspace-sts-5.3.0\mafiagame")
WORK = ROOT / ".docx_work"
REFERENCE = Path(r"C:\Users\KOSA_L3\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-system-design\assets\reference.docx")
FINAL = ROOT / "MAFIAGAME_MVC_구조_데이터_흐름.docx"
ARCH_IMAGE = WORK / "mafiagame-mvc-architecture.png"
STAGED = WORK / "mafiagame-mvc-staged.docx"


NAVY = "#0B2F51"
BLUE = "#3A78A6"
PALE = "#E7F1F9"
PALE2 = "#F3F8FC"
TEXT = "#23394C"
MUTED = "#5D7489"
WHITE = "#FFFFFF"


def font(path: str, size: int):
    return ImageFont.truetype(path, size=size)


def make_architecture_image(path: Path) -> None:
    width, height = 1800, 920
    img = Image.new("RGB", (width, height), "#F3F8FC")
    draw = ImageDraw.Draw(img)
    regular_path = r"C:\Windows\Fonts\malgun.ttf"
    bold_path = r"C:\Windows\Fonts\malgunbd.ttf"
    regular = font(regular_path, 25)
    small = font(regular_path, 21)
    tiny = font(regular_path, 18)
    bold = font(bold_path, 30)
    title = font(bold_path, 38)

    draw.rectangle((0, 0, width, 126), fill=NAVY)
    draw.text((56, 27), "MAFIAGAME MVC 연결 구조", fill=WHITE, font=title)
    draw.text((58, 79), "HTTP 화면 요청과 WebSocket 실시간 이벤트가 같은 서비스 계층을 공유", fill="#DCEAF5", font=small)

    def box(x, y, w, h, heading, lines, fill=WHITE, outline=NAVY):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=fill, outline=outline, width=4)
        draw.text((x + 24, y + 20), heading, fill=NAVY, font=bold)
        for idx, line in enumerate(lines):
            draw.text((x + 24, y + 77 + idx * 34), line, fill=TEXT, font=small)

    def arrow(x1, y1, x2, y2, label=None, dashed=False):
        draw.line((x1, y1, x2, y2), fill=BLUE, width=5)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        size = 19
        left = (x2 - size * math.cos(angle - 0.45), y2 - size * math.sin(angle - 0.45))
        right = (x2 - size * math.cos(angle + 0.45), y2 - size * math.sin(angle + 0.45))
        draw.polygon([(x2, y2), left, right], fill=BLUE)
        if label:
            midx, midy = (x1 + x2) // 2, (y1 + y2) // 2
            draw.text((midx - 45, midy - 34), label, fill=MUTED, font=tiny)

    top_y, box_w, box_h, gap = 208, 286, 172, 30
    xs = [62, 378, 694, 1010, 1326]
    box(xs[0], top_y, box_w, box_h, "View", ["Thymeleaf 화면", "chat.js DOM 갱신"])
    box(xs[1], top_y, box_w, box_h, "Controller", ["Auth · Room · User", "Chat · Presence"])
    box(xs[2], top_y, box_w, box_h, "Service", ["입력 검증과 업무 규칙", "트랜잭션 · 실시간 상태"])
    box(xs[3], top_y, box_w, box_h, "Mapper", ["Mapper 인터페이스", "MyBatis XML SQL"])
    box(xs[4], top_y, box_w, box_h, "MariaDB", ["user", "game_room · room_members"])
    for i in range(4):
        arrow(xs[i] + box_w, top_y + box_h // 2, xs[i + 1], top_y + box_h // 2)

    lower_y = 584
    box(146, lower_y, 410, 160, "SecurityConfig", ["폼 로그인 · BCrypt", "인증된 /ws 접근만 허용"], fill=PALE)
    box(696, lower_y, 410, 160, "WebSocketConfig", ["/ws 엔드포인트", "/app → /topic · /queue"], fill=PALE)
    box(1246, lower_y, 410, 160, "실시간 상태", ["RoomPresenceService", "roomId · sessionId Map"], fill=PALE)
    arrow(xs[1] + box_w // 2, top_y + box_h, 356, lower_y, "보호")
    arrow(xs[2] + box_w // 2, top_y + box_h, 901, lower_y, "메시지")
    arrow(xs[2] + box_w, top_y + 123, 1246, lower_y + 80, "상태")
    draw.text((58, 845), "실선은 요청 또는 데이터 전달, 하단 영역은 인증과 실시간 기능의 보조 경로", fill=MUTED, font=tiny)
    img.save(path)


def set_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def set_plain_text(paragraph, text: str) -> None:
    """Replace a paragraph including runs nested inside a hyperlink."""
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    paragraph.add_run(text)


def set_cell(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    set_text(cell.paragraphs[0], text)
    for paragraph in cell.paragraphs[1:]:
        set_text(paragraph, "")


def replace_paragraphs(doc: Document, replacements: dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text in replacements:
            set_text(paragraph, replacements[paragraph.text])


def replace_tables(doc: Document) -> None:
    tables = doc.tables

    # Title-page status row.
    set_cell(tables[0].cell(0, 0), "상태\n구현 기준")
    set_cell(tables[0].cell(0, 2), "소유자\n개인 프로젝트")
    set_cell(tables[0].cell(0, 4), "최종 수정\n2026년 9월 16일")

    metadata = [
        ("작성자", "개인 프로젝트 개발자"),
        ("검토자", "본인 검토용"),
        ("관련 문서", "README.md · mafiasql.sql"),
        ("범위", "현재 구현된 Spring MVC, Thymeleaf, MyBatis, MariaDB, WebSocket/STOMP 연결과 데이터 흐름"),
    ]
    for r, (left, right) in enumerate(metadata):
        set_cell(tables[1].cell(r, 0), left)
        set_cell(tables[1].cell(r, 1), right)

    goals = [
        ("목표", "제외 범위"),
        ("Controller를 HTTP와 STOMP 메시지의 진입점으로 유지", "직업 배정, 낮과 밤, 투표, 승패 판정까지의 전체 게임 진행"),
        ("Service에서 정규화, 검증, 해시, 트랜잭션을 처리", "채팅 메시지와 게임 전적의 영구 저장"),
        ("Mapper와 XML SQL을 분리해 MariaDB 접근을 한 곳에서 관리", "여러 서버에서 공유되는 분산 실시간 상태"),
        ("대기실 참가자와 준비 상태를 WebSocket 스냅샷으로 동기화", "방 입장 비밀번호, 정원, 중복 입장의 서버 검증은 현재 미완성"),
    ]
    for r, row in enumerate(goals):
        for c, value in enumerate(row):
            set_cell(tables[2].cell(r, c), value)

    components = [
        ("구성 요소", "책임", "주요 저장소", "실패 동작"),
        ("View 계층", "Thymeleaf 템플릿이 서버 모델을 HTML로 렌더링하고 chat.js가 채팅과 참가자 DOM을 갱신", "브라우저 DOM", "화면은 서버 응답 또는 WebSocket 스냅샷이 없으면 최신 상태를 표시하지 못함"),
        ("Controller 계층", "AuthController와 RoomController가 HTTP를 처리하고 ChatController와 RoomPresenceController가 STOMP 메시지를 받음", "HTTP 세션과 메시징 채널", "입력 오류와 RoomWebSocketException을 화면 또는 사용자 큐 오류로 반환"),
        ("Service 계층", "SignupService와 RoomService가 검증·정규화·암호화·트랜잭션을 담당하고 RoomPresenceService가 실시간 상태를 관리", "MariaDB와 인메모리 Map", "규칙 위반은 작업을 거부하고 DB 오류는 트랜잭션 경계에서 실패"),
        ("Persistence 계층", "Mapper 인터페이스와 XML이 사용자·방·멤버 조회 및 저장 SQL을 실행", "MariaDB user, game_room, room_members", "SQL 오류나 제약 조건 위반이 상위 계층으로 전달"),
        ("Security와 transport", "SecurityConfig가 폼 로그인과 경로 접근을 보호하고 WebSocketConfig가 /ws와 STOMP 목적지를 등록", "JSESSIONID와 단순 브로커", "인증되지 않은 일반 요청과 WebSocket 연결을 허용하지 않음"),
    ]
    for r, row in enumerate(components):
        for c, value in enumerate(row):
            set_cell(tables[3].cell(r, c), value)

    contracts = [
        ("계약", "형식", "필수", "의미와 처리 규칙"),
        ("RoomList", "DB 조회 모델", "예", "방 목록과 상세 화면에 쓰는 방 ID, 제목, 방장, 현재 인원, 최대 인원, 상태, 잠금 여부"),
        ("RoomView", "서버 view model", "예", "rooms/detail 템플릿에 전달하는 제목, 설명, 인원 수, 정원, 상태, 잠금 여부"),
        ("ChatMessageRequest.content", "String", "예", "STOMP SEND 입력. strip 후 공백을 거부하고 코드 포인트 기준 300자까지 허용"),
        ("ChatMessage", "roomId, type, sender, content, sentAt", "예", "검증된 채팅을 /topic/rooms/{roomId}/chat으로 브로드캐스트. type은 CHAT"),
        ("RoomParticipant", "userId, nickname, host, ready", "예", "현재 WebSocket 세션의 참가자 표현. 방장과 준비 상태를 포함"),
        ("RoomPresenceState", "roomId, participants[]", "예", "참가자 전체 스냅샷을 /topic/rooms/{roomId}/presence로 전송"),
        ("ChatError", "type, message", "조건부", "처리 오류를 /user/queue/errors로 개인 전송. type은 ERROR"),
    ]
    for r, row in enumerate(contracts):
        for c, value in enumerate(row):
            set_cell(tables[4].cell(r, c), value)

    consistency = [
        ("시나리오", "기대 동작", "근거"),
        ("같은 세션의 재연결", "기존 세션을 방에서 제거하고 새 참가자 스냅샷을 전송", "roomBySession과 participantsByRoom을 함께 갱신해 중복 세션을 줄임"),
        ("DB 저장 또는 제약 조건 실패", "회원가입·방 생성이 예외로 종료되고 사용자에게 입력 또는 재시도 안내", "SignupService와 RoomService가 @Transactional 경계를 사용"),
        ("WebSocket 연결 종료", "세션을 인메모리 Map에서 제거하고 남은 참가자에게 스냅샷 전송", "SessionDisconnectEvent를 RoomPresenceService가 수신"),
        ("동일 채팅 요청 반복", "요청마다 CHAT 메시지가 생성될 수 있음", "현재 ChatMessage에 idempotency key가 없으므로 중복 전송 방지는 미구현"),
    ]
    for r, row in enumerate(consistency):
        for c, value in enumerate(row):
            set_cell(tables[5].cell(r, c), value)

    operational = [
        ("신호", "현재 확인 기준", "담당", "출시 기준"),
        ("HTTP 화면 응답", "뷰 이름 또는 redirect 반환, Spring 컨텍스트 테스트 통과", "개인 개발자", "필수"),
        ("WebSocket 연결과 재연결", "chat.js가 연결 상태를 표시하고 3초 후 재연결", "개인 개발자", "필수"),
        ("DB 조회와 트랜잭션", "Mapper SQL과 서비스 예외 메시지로 실패 확인", "개인 개발자", "필수"),
        ("Presence 정합성", "join, ready, disconnect 후 전체 참가자 스냅샷 브로드캐스트", "개인 개발자", "필수"),
        ("보안 경계", "/ws와 보호 경로에 인증 필요, 비밀번호는 BCrypt 해시 저장", "개인 개발자", "필수"),
        ("출시 제약", "단일 서버·인메모리 presence 전제. 방 입장 검증, 게임 상태, 통합 테스트, 운영 로그를 보강한 뒤 공개", "개인 개발자", "권장"),
    ]
    for r, row in enumerate(operational):
        for c, value in enumerate(row):
            set_cell(tables[6].cell(r, c), value)

    alternatives = [
        ("대안", "검토 이유", "선택하지 않은 이유"),
        ("SPA + REST API", "화면과 API를 분리하면 풍부한 클라이언트 상태 관리에 유리", "현재 MVP 화면은 Thymeleaf로 충분하고 프론트엔드 복잡도를 늘리지 않음"),
        ("JPA 엔티티 중심 접근", "객체 관계와 CRUD를 빠르게 모델링할 수 있음", "현재 SQL 집계와 스키마 확인이 중요한 방 목록에는 MyBatis XML이 명시적"),
        ("Redis 기반 presence", "여러 서버에서 세션 상태를 공유할 수 있음", "현재는 단일 서버 범위이며 Redis 운영 부담은 후속 확장 과제로 남김"),
        ("SockJS fallback", "브라우저 호환성과 연결 대체 경로를 제공", "동일 출처의 기본 WebSocket을 사용하기로 했고 별도 fallback 요구가 없음"),
    ]
    for r, row in enumerate(alternatives):
        for c, value in enumerate(row):
            set_cell(tables[7].cell(r, c), value)

    milestones = [
        ("단계", "산출물", "완료 기준"),
        ("M1", "방 입장·퇴장·비밀번호·정원 검증", "서버가 방 상태와 사용자 권한을 기준으로 허용·거부하고 테스트로 확인"),
        ("M2", "게임 상태 모델과 단계 전이 테스트", "WAITING부터 RESULT까지의 상태 전이와 잘못된 행동을 검증"),
        ("M3", "단일 서버 운영 준비", "인증, WebSocket, DB, 재접속, 오류 경로의 통합 테스트와 기본 로그 확보"),
        ("M4", "확장 설계 검토", "동시 방 운영과 다중 서버가 필요할 때 Redis Pub/Sub 및 분산 상태 저장을 결정"),
    ]
    for r, row in enumerate(milestones):
        for c, value in enumerate(row):
            set_cell(tables[8].cell(r, c), value)


def replace_image_in_package(docx_path: Path, image_path: Path) -> None:
    temp = docx_path.with_suffix(".imagepatched.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = image_path.read_bytes() if item.filename == "word/media/image1.png" else zin.read(item.filename)
            zout.writestr(item, data)
    temp.replace(docx_path)


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    make_architecture_image(ARCH_IMAGE)
    shutil.copyfile(REFERENCE, STAGED)
    doc = Document(STAGED)

    replacements = {
        "System Name": "MAFIAGAME",
        "Title of Proposal": "MVC 패턴 연결 구조와 데이터 처리 흐름",
        "1.  Abstract": "1 개요",
        "2.  Goals and Non-Goals": "2 목표와 제외 범위",
        "3.  Background and Problem Statement": "3 배경과 문제 정의",
        "4.  Proposed Architecture": "4 현재 아키텍처",
        "5.  Request Lifecycle": "5 요청과 이벤트 처리 흐름",
        "6.  API and Data Contracts": "6 API와 데이터 계약",
        "7.  Consistency, Idempotency, and Replay": "7 일관성과 재연결 처리",
        "7. Consistency, Idempotency, and Replay": "7 일관성과 재연결 처리",
        "8.  Security and Privacy Considerations": "8 보안과 개인정보 고려사항",
        "8. Security and Privacy Considerations": "8 보안과 개인정보 고려사항",
        "9.  Operational Readiness": "9 운영 준비 상태",
        "10. Alternatives Considered": "10 대안 비교",
        "11. Open Questions": "11 남은 질문",
        "12. Decision and Next Steps": "12 결정과 다음 단계",
        "Core components": "핵심 구성 요소",
        "Primary data contract": "핵심 런타임 계약",
        "Contract guarantees": "계약 보장 사항",
        "The versioned schema or interface definition is published at ": "버전 기준은 다음 소스에 반영한다",
        "[Link to interface or schema] and update it with each contract release.": "README.md, mafiasql.sql, Mapper XML을 변경 시 함께 검토한다.",
        "[Summarize the proposed system, the problem it solves, and the intended outcome. Describe the core design at a high level, including the main boundaries, dependencies, and guarantees. Keep this section concise enough that a reviewer can understand the proposal without reading the full document.]": "이 문서는 MAFIAGAME의 현재 구현을 기준으로 MVC 계층과 실시간 대기실 기능이 어떻게 연결되는지 설명한다. 브라우저의 Thymeleaf 화면과 chat.js는 HTTP 요청 또는 WebSocket 이벤트를 보내고, Controller가 요청 경계를 받는다. Service는 입력을 정규화하고 업무 규칙을 적용하며, Mapper와 MyBatis XML은 MariaDB의 사용자·방·멤버 데이터를 처리한다. 실시간 참가자와 준비 상태는 RoomPresenceService의 인메모리 Map에서 관리하고 STOMP topic으로 전체 스냅샷을 전파한다.",
        "[Describe the workloads and constraints this design must support. Clarify what the proposal does not attempt to solve, the assumptions it relies on, and the most important operational or implementation boundaries.]": "현재 범위는 회원가입·로그인, 방 목록·생성·상세, 인증된 사용자의 대기방 채팅, 참가자 표시, 준비 상태 동기화다. 회원·방 메타데이터는 DB가 기준이고 실시간 접속 상태는 단일 서버의 메모리가 기준이다. 게임의 직업 배정, 낮과 밤, 투표, 승패 판정과 서버 측 방 입장 검증은 아직 후속 구현 범위다.",
        "[Describe the current state, the specific problem, and why the existing approach is no longer sufficient. Include relevant scale, reliability, security, cost, or developer-experience constraints, and explain the impact of leaving the problem unresolved.]": "프로젝트는 Spring Boot 기반의 단일 애플리케이션으로 화면 렌더링, 인증, DB 접근, 실시간 대기실을 함께 제공한다. 기능이 늘어날수록 화면 모델을 만드는 Controller, 검증과 규칙을 담는 Service, SQL을 실행하는 Mapper의 경계가 흐려질 수 있다. 연결 구조를 명시하지 않으면 게임 상태 로직을 잘못된 계층에 넣거나 DB 상태와 WebSocket 상태를 혼동하기 쉽다.",
        "[Explain the proposed system boundary and the major responsibilities on each side. Name the primary components, how they interact, and which inputs determine behavior. State the key design principle or invariant that should guide implementation and review.]": "외부 경계는 브라우저와 인증된 WebSocket 클라이언트이며 내부 경계는 Controller, Service, Persistence로 나뉜다. 요청은 반드시 Controller를 거쳐 Service의 규칙을 통과하고, 영구 데이터 변경은 Mapper를 통해서만 수행한다. WebSocket 상태는 sessionId를 키로 방별 참가자 스냅샷을 만들고, 연결 종료 시 같은 경로로 제거한다. 핵심 원칙은 화면이나 메시지 핸들러가 DB와 규칙을 직접 조작하지 않는 것이다.",
        "Figure 1. [Proposed System Architecture].": "그림 1. MAFIAGAME MVC 연결 구조",
        "[Describe how a request, event, or job enters the system and identify the required inputs.]": "HTTP는 브라우저가 /login, /signup, /rooms, /rooms/new, /rooms/{roomId}로 요청하면서 시작한다. 방 페이지의 chat.js는 같은 출처의 /ws에 WebSocket을 열고 STOMP CONNECT 후 방별 topic을 구독한다. 채팅과 준비 변경은 /app/rooms/{roomId}/chat 또는 /ready로 SEND하며, 입장 알림은 /join으로 보낸다.",
        "[Describe validation, authentication, authorization, and normalization at the system boundary.]": "Spring Security 필터가 경로와 인증 여부를 먼저 확인한다. 로그인은 email과 password를 사용하고, /ws/**는 인증된 사용자만 접근한다. Controller가 받은 값은 Service가 trim, 이메일 소문자화, 길이·범위·동의 여부 검사를 수행한다. 채팅은 strip 후 공백과 300자 초과를 거부한다.",
        "[Describe which configuration, policy, state, or dependency data is loaded before processing.]": "회원가입과 로그인은 UserMapper로 사용자 정보를 읽고 PasswordEncoder로 비밀번호를 비교한다. 방 목록과 상세는 RoomMapper가 방장 이름과 room_members 집계를 포함해 조회한다. Presence join은 방이 존재하는지 RoomService로 확인한 뒤 인증 주체의 userId와 nickname을 세션 상태에 넣는다. 브로커 경로와 /ws 엔드포인트는 WebSocketConfig가 제공한다.",
        "[Describe the primary decision or processing step and the output it produces.]": "SignupService는 입력을 검증하고 BCrypt 해시를 만든 뒤 UserMapper.insert를 실행한다. RoomService는 제목·정원·비밀번호를 검증하고 방을 만든 뒤 방장을 room_members에 추가한다. ChatController는 메시지를 검증해 ChatMessage를 만들고 room topic으로 전송한다. RoomPresenceService는 ready 값을 갱신해 RoomPresenceState 전체 스냅샷을 전송한다.",
        "[Describe which durable state must be written before side effects or downstream execution begin.]": "회원가입은 사용자 행이 저장되어야 로그인 대상이 된다. 방 생성은 game_room 행과 방장 멤버 행이 함께 성공해야 완료되며 두 작업은 하나의 트랜잭션 안에서 수행된다. 반면 채팅과 presence, ready는 현재 DB에 저장하지 않고 브로커와 인메모리 상태로만 전달한다.",
        "[Describe downstream calls, retries, timeouts, budget limits, and terminal conditions.]": "HTTP 요청은 Mapper SQL을 실행한 뒤 view 이름 또는 redirect를 반환한다. WebSocket 메시지는 Spring Simple Broker의 /topic, /queue로 전달된다. 브라우저는 연결 종료 후 3초 뒤 재연결하며, 연결 전에는 채팅 전송과 준비 버튼을 비활성화한다. 현재 DB 재시도 정책이나 메시지 idempotency key는 없으며 오류는 사용자별 /user/queue/errors로 전달된다.",
        "[Describe final state updates, the response or output, and the metrics, logs, and traces emitted.]": "성공한 HTTP 요청은 rooms/list, rooms/create, rooms/detail 또는 로그인 redirect로 마무리된다. 채팅은 메시지 목록에 sender·content·sentAt을 추가하고, presence는 참가자 카드와 현재 인원을 다시 그린다. disconnect는 Map에서 세션을 제거한 뒤 남은 참가자에게 상태를 브로드캐스트한다. 현재 운영 메트릭과 분산 추적은 별도 구현 전이다.",
        "[State the ordering, durability, or validation guarantee that must hold before execution.]": "정규화와 검증을 통과한 입력만 Service의 저장 또는 브로드캐스트 단계로 넘어간다. 방 생성은 room_id 생성 후 방장 멤버 등록까지 성공해야 정상 종료하며, 회원가입과 방 생성은 @Transactional로 묶인다.",
        "[State how writes, attempts, or events are identified and ordered.]": "DB 행은 user_id와 room_id로 식별하고 room_members는 room_id·user_id 복합 키를 사용한다. 실시간 이벤트는 WebSocket sessionId와 roomId로 식별하며, 참가자 목록은 방장 우선·닉네임·userId 순으로 정렬한다. 채팅에는 이벤트 ID나 순서 보장 필드가 없다.",
        "[State which versions, identifiers, or inputs must be captured for audit and replay.]": "화면과 메시지 처리에는 roomId, userId 또는 sessionId, 현재 ready 상태가 필요하다. 채팅은 sender와 sentAt을 함께 보내지만 서버 로그나 DB에 저장하지 않는다. 운영 감사와 재생이 필요해지면 이벤트 ID, 연결 세션, 서버 시각, 게임 단계 저장을 추가해야 한다.",
        "[State what this data is and is not a source of truth for.]": "MariaDB는 사용자 계정, 방 메타데이터, 방 멤버의 영구 기준이다. RoomPresenceService의 Map은 현재 연결된 브라우저와 준비 상태의 기준이지만 프로세스 재시작 후 복원되지 않는다. HTML의 초기 참가자 수는 렌더링 시점 값이며 WebSocket 스냅샷이 도착하면 갱신된다.",
        "[Explain the consistency, idempotency, replay, ordering, or concurrency guarantees required by this design. Distinguish client-facing guarantees from internal analysis or recovery behavior, and identify where duplicate work or partial failure is acceptable.]": "현재 구현은 DB 영구 상태와 실시간 메모리 상태를 분리한다. synchronized monitor가 한 프로세스 안에서 참가자 Map 변경과 스냅샷 생성을 보호하고, disconnect와 재연결 때 방 전체에 새 스냅샷을 전송한다. 하지만 채팅은 중복 요청을 판별하지 않으며 다중 서버 간 상태 공유도 보장하지 않는다.",
        "[Describe authentication, authorization, and tenant or data-boundary requirements.]": "일반 화면 중 /, /rooms, /login, /signup과 정적 리소스는 허용되고, 그 외 경로는 인증이 필요하다. 폼 로그인은 email을 사용자 식별자로 사용하며, /ws/** 연결도 인증된 사용자만 허용한다. WebSocket 메시지의 방 참가 여부 검증은 join과 ready 경로의 현재 검사 범위 안에서만 수행된다.",
        "[Describe data minimization, sensitive payload handling, and logging restrictions.]": "비밀번호는 BCrypt로 저장하고 User의 toString에서 제외한다. 채팅 DOM 출력은 textContent를 사용해 HTML로 해석되지 않도록 한다. 현재 채팅과 게임 전적은 저장하지 않으므로 보존 데이터는 최소화되어 있지만, 운영 로그를 추가할 때 비밀번호와 민감한 인증 정보를 기록하지 않아야 한다.",
        "[Describe credential, secret, key, and certificate storage and access requirements.]": "DB 사용자와 비밀번호는 application.properties의 DB_USERNAME, DB_PASSWORD 환경 변수로 주입한다. 애플리케이션은 PasswordEncoder Bean으로 BCrypt를 사용하며 평문 비밀번호를 DB에 저장하지 않는다. 운영에서는 환경 변수 또는 비밀 저장소를 사용하고 소스와 로그에 자격 증명을 남기지 않아야 한다.",
        "[Describe safe defaults for administrative, replay, migration, and debugging tools.]": "인증되지 않은 /ws 연결을 기본 거부하고, presence 확인 전에는 준비 버튼을 비활성화한다. 잘못된 메시지와 존재하지 않는 방은 RoomWebSocketException으로 거부한다. 다만 상세 화면의 존재하지 않는 방 mock fallback과 사용자 프로필의 샘플 데이터는 개발용 동작이므로 운영 전 제거하거나 명확한 오류 화면으로 바꿔야 한다.",
        "[Describe retention, deletion, residency, privacy, and audit requirements.]": "현재 스키마는 사용자, 방, 방 멤버를 저장하고 채팅과 게임 기록은 저장하지 않는다. 탈퇴·방 삭제·멤버 정리의 보존 정책과 외래 키 삭제 방식은 아직 결정되지 않았다. 전적과 게임 기록을 도입할 때는 보존 기간, 삭제 요청, 감사 이벤트를 함께 정의해야 한다.",
        "Rollout constraint: [Describe dry-run, canary, feature-flag, validation, rollback, and promotion requirements before production launch.]": "출시는 단일 서버와 인메모리 presence를 전제로 한다. 공개 전 방 입장·퇴장·비밀번호·정원 검증, 존재하지 않는 방 처리, WebSocket 권한 검증, 상태 전이 테스트, 오류 로그와 재접속 테스트를 완료한다. 서버 교체 시 presence가 초기화된다는 제약을 운영 문서에 명시한다.",
        "[State the recommended decision and summarize the implementation sequence. Name the first milestone, the validation or dry-run stage, the initial production audience, and the conditions that must be met before broader rollout.]": "현재 MVP에서는 Spring MVC와 Thymeleaf를 View 경계로 유지하고, Controller → Service → Mapper 구조를 기준선으로 삼는다. WebSocket 채팅과 presence는 Controller에서 메시지를 받고 Service가 상태 변경과 브로드캐스트를 담당하게 한다. 첫 단계는 방 입장 검증과 게임 상태 모델을 완성하는 것이며, 단일 서버 테스트와 운영 로그를 통과한 뒤 다중 서버 확장 필요성을 재평가한다.",
    }
    replace_paragraphs(doc, replacements)

    # Some template runs contain a different whitespace split after import.
    # Resolve those slots by semantic marker so no template text survives.
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if "Consistency, Idempotency" in text:
            set_text(paragraph, "7 일관성과 재연결 처리")
        elif "Security and Privacy" in text:
            set_text(paragraph, "8 보안과 개인정보 고려사항")
        elif "Link to interface or schema" in text:
            set_plain_text(paragraph, "README.md, mafiasql.sql, Mapper XML을 변경 시 함께 검토한다.")
        elif "Open question 1:" in text:
            set_text(paragraph, "방 입장·퇴장·비밀번호·정원 검증을 어느 시점에 완료할 것인가?")
        elif "Open question 2:" in text:
            set_text(paragraph, "채팅과 게임 전적을 저장할지, 저장한다면 보존 기간은 얼마인가?")
        elif "Open question 3:" in text:
            set_text(paragraph, "단일 서버의 인메모리 상태로 충분한 동시 방 규모는 얼마인가?")
        elif "Open question 4:" in text:
            set_text(paragraph, "게임 상태 전이와 역할별 권한의 최종 소유 Service는 무엇인가?")
    replace_tables(doc)

    # Update the recurring footer while preserving its source formatting.
    footer = doc.sections[0].footer
    if footer.paragraphs:
        set_text(footer.paragraphs[0], "MAFIAGAME | MVC 구조 문서")

    if doc.inline_shapes:
        doc.inline_shapes[0]._inline.docPr.set("descr", "MAFIAGAME MVC 연결 구조와 데이터 처리 흐름")
        doc.inline_shapes[0]._inline.docPr.set("title", "MAFIAGAME MVC 연결 구조")

    doc.save(STAGED)
    replace_image_in_package(STAGED, ARCH_IMAGE)
    STAGED.replace(FINAL)


if __name__ == "__main__":
    main()
